import io
import re
from collections import defaultdict
import pdfplumber
from fastapi import APIRouter, Depends, File, HTTPException, Query, status, Response, UploadFile
from fastapi.responses import StreamingResponse
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter
from sqlalchemy.orm import Session
from typing import List, Optional
from datetime import datetime
from ..database import get_db
from ..models.academic import Student, Subject, Faculty, FacultySubject
from ..models.performance import Marks, Attendance, SemesterResult, UniversityExamResult
from ..models.user import User
from ..schemas.performance import MarksheetOut, MarksheetSubjectItem
from ..schemas.results import SpreadsheetExportRequest, UniversityResultImportRequest
from ..services.auth import get_current_user, require_role
from ..services.pdf_service import generate_marksheet_pdf
from ..services.grading import recalculate_student_results
from ..services.audit_service import record_audit_log

router = APIRouter(prefix="/api/results", tags=["Results & Marksheets"])


@router.post("/export-excel")
def export_results_to_excel(
    payload: SpreadsheetExportRequest,
    current_user: User = Depends(require_role(["admin", "faculty"]))
):
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "Academic Report"
    worksheet.freeze_panes = "A2"
    worksheet.auto_filter.ref = f"A1:{get_column_letter(len(payload.columns))}{len(payload.rows) + 1}"
    worksheet.append([column.header for column in payload.columns])

    for cell in worksheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill(fill_type="solid", fgColor="1E3A8A")

    for row in payload.rows:
        values = []
        for column in payload.columns:
            value = row.get(column.key)
            if isinstance(value, (dict, list, tuple)):
                value = str(value)
            if isinstance(value, str) and value.startswith(("=", "+", "-", "@")):
                value = f"'{value}"
            values.append(value)
        worksheet.append(values)

    for column_cells in worksheet.columns:
        column_letter = get_column_letter(column_cells[0].column)
        max_length = max(len(str(cell.value or "")) for cell in column_cells)
        worksheet.column_dimensions[column_letter].width = min(max(max_length + 2, 12), 48)

    output = io.BytesIO()
    workbook.save(output)
    output.seek(0)
    filename = re.sub(r"[^a-zA-Z0-9_-]+", "_", payload.title).strip("_") or "academic_report"
    return StreamingResponse(
        output,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{filename}.xlsx"'}
    )


@router.post("/extract-pdf-text")
async def extract_university_pdf_text(
    file: UploadFile = File(...),
    current_user: User = Depends(require_role(["admin", "faculty"]))
):
    if not (file.filename or "").lower().endswith(".pdf"):
        raise HTTPException(status_code=400, detail="Upload a PDF result sheet")
    contents = await file.read()
    try:
        with pdfplumber.open(io.BytesIO(contents)) as document:
            pages = [page.extract_text() or "" for page in document.pages]
            text = "\n".join(pages).strip()
            used_ocr = False
            if not text:
                try:
                    import pytesseract
                except ImportError as error:
                    raise HTTPException(
                        status_code=503,
                        detail="Scanned PDFs require the pytesseract package and Tesseract OCR engine."
                    ) from error
                try:
                    pages = [
                        pytesseract.image_to_string(page.to_image(resolution=250).original)
                        for page in document.pages
                    ]
                    text = "\n".join(pages).strip()
                    used_ocr = True
                except pytesseract.TesseractNotFoundError as error:
                    raise HTTPException(
                        status_code=503,
                        detail="This scanned PDF needs the Tesseract OCR engine installed on the API host."
                    ) from error
                except Exception as error:
                    raise HTTPException(status_code=422, detail=f"OCR could not read this PDF: {error}") from error
    except HTTPException:
        raise
    except Exception as error:
        raise HTTPException(status_code=400, detail=f"Could not read the PDF: {error}") from error

    if not text:
        raise HTTPException(status_code=422, detail="No readable text was found in this PDF, including the OCR pass.")
    return {"text": text, "used_ocr": used_ocr}


def construct_marksheet_dict(student_id: int, semester: int, db: Session) -> dict:
    student = db.query(Student).filter(Student.id == student_id).first()
    if not student:
        raise HTTPException(status_code=404, detail="Student not found")

    marks = db.query(Marks).filter(
        Marks.student_id == student_id,
        Marks.semester == semester
    ).all()

    if not marks:
        # Check if student has any marks in this semester
        pass

    subjects_list = []
    total_creds = 0
    earned_creds = 0
    has_failed = False

    for m in marks:
        subj = m.subject
        c = subj.credits if subj else 4
        total_creds += c
        cp = round(c * m.grade_point, 1)

        if m.is_passed:
            earned_creds += c
        else:
            has_failed = True

        subjects_list.append({
            "subject_code": subj.code if subj else "N/A",
            "subject_name": subj.name if subj else "N/A",
            "credits": c,
            "max_internal": subj.max_internal if subj else 30,
            "max_external": subj.max_external if subj else 70,
            "internal_marks": m.internal_marks,
            "external_marks": m.external_marks,
            "total_marks": m.total_marks,
            "grade": m.grade,
            "grade_point": m.grade_point,
            "credit_points": cp,
            "is_passed": m.is_passed,
            "result": "PASS" if m.is_passed else "FAIL"
        })

    # Fetch SemesterResult record
    sem_res = db.query(SemesterResult).filter(
        SemesterResult.student_id == student_id,
        SemesterResult.semester == semester
    ).first()

    if not sem_res:
        recalculate_student_results(student_id, db)
        sem_res = db.query(SemesterResult).filter(
            SemesterResult.student_id == student_id,
            SemesterResult.semester == semester
        ).first()

    sgpa = sem_res.sgpa if sem_res else 0.0
    cgpa = sem_res.cgpa if sem_res else 0.0
    status_str = "FAIL" if has_failed else "PASS"

    # Classification
    if status_str == "FAIL":
        classification = "Reappear / Backlog Required"
    elif cgpa >= 8.5:
        classification = "First Class with Distinction (Exemplary)"
    elif cgpa >= 6.5:
        classification = "First Class"
    elif cgpa >= 5.0:
        classification = "Second Class"
    else:
        classification = "Pass Class"

    # Attendance
    att_recs = db.query(Attendance).filter(
        Attendance.student_id == student_id,
        Attendance.semester == semester
    ).all()
    if att_recs:
        held = sum(a.classes_held for a in att_recs)
        att = sum(a.classes_attended for a in att_recs)
        att_pct = round((att / held * 100), 1) if held > 0 else 85.0
    else:
        att_pct = 85.0

    return {
        "student_id": student.id,
        "register_number": student.register_number,
        "student_name": student.full_name,
        "department_name": student.department.name if student.department else "Engineering",
        "department_code": student.department.code if student.department else "ENG",
        "year": student.year,
        "semester": semester,
        "section": student.section,
        "email": student.email,
        "subjects": subjects_list,
        "total_credits": total_creds or 20,
        "earned_credits": earned_creds,
        "sgpa": sgpa,
        "cgpa": cgpa,
        "result_status": status_str,
        "classification": classification,
        "attendance_percentage": att_pct,
        "issue_date": datetime.utcnow().strftime("%B %d, %Y")
    }

@router.get("/marksheet/{student_id}/{semester}", response_model=MarksheetOut)
def get_official_marksheet(
    student_id: int,
    semester: int,
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user)
):
    # Security: student can only view their own
    if current_user.role == "student":
        stud = db.query(Student).filter(Student.user_id == current_user.id).first()
        if not stud or stud.id != student_id:
            raise HTTPException(status_code=403, detail="Unauthorized to view this marksheet")

    data = construct_marksheet_dict(student_id, semester, db)
    return data

@router.get("/marksheet/{student_id}/{semester}/pdf")
def download_marksheet_pdf(
    student_id: int,
    semester: int,
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user)
):
    if current_user.role == "student":
        stud = db.query(Student).filter(Student.user_id == current_user.id).first()
        if not stud or stud.id != student_id:
            raise HTTPException(status_code=403, detail="Unauthorized to download this marksheet")

    data = construct_marksheet_dict(student_id, semester, db)
    pdf_buffer = generate_marksheet_pdf(data)

    reg = data["register_number"]
    filename = f"Marksheet_{reg}_Sem{semester}.pdf"

    return Response(
        content=pdf_buffer.getvalue(),
        media_type="application/pdf",
        headers={"Content-Disposition": f"attachment; filename={filename}"}
    )

@router.get("/student/{student_id}")
def get_all_results_for_student(
    student_id: int,
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user)
):
    if current_user.role == "student":
        stud = db.query(Student).filter(Student.user_id == current_user.id).first()
        if not stud or stud.id != student_id:
            raise HTTPException(status_code=403, detail="Unauthorized to view these semester results")

    res = db.query(SemesterResult).filter(SemesterResult.student_id == student_id).order_by(SemesterResult.semester.asc()).all()
    return res


@router.post("/university/import")
def import_university_results(
    payload: UniversityResultImportRequest,
    db: Session = Depends(get_db),
    current_user: User = Depends(require_role(["admin", "faculty"]))
):
    normalized_rows = []
    seen_keys = set()
    faculty = None
    if current_user.role == "faculty":
        faculty = db.query(Faculty).filter(Faculty.user_id == current_user.id).first()
        if not faculty:
            raise HTTPException(status_code=403, detail="Faculty profile is required to import university results")

    for row in payload.records:
        register_number = row.register_number.strip().upper()
        subject_code = row.subject_code.strip().upper()
        exam_session = row.exam_session.strip()
        if not exam_session:
            raise HTTPException(status_code=400, detail="Examination session cannot be blank")
        key = (register_number, subject_code, row.semester, exam_session)
        if key in seen_keys:
            raise HTTPException(
                status_code=400,
                detail=f"Duplicate result row for {register_number} / {subject_code} in {row.exam_session}"
            )
        seen_keys.add(key)

        student = db.query(Student).filter(Student.register_number == register_number).first()
        if not student:
            raise HTTPException(status_code=400, detail=f"Student '{register_number}' was not found")
        subject = db.query(Subject).filter(Subject.code == subject_code).first()
        if not subject:
            raise HTTPException(status_code=400, detail=f"Subject '{subject_code}' was not found")

        if faculty:
            assignment = db.query(FacultySubject).filter(
                FacultySubject.faculty_id == faculty.id,
                FacultySubject.subject_id == subject.id
            ).first()
            if not assignment:
                raise HTTPException(
                    status_code=403,
                    detail=f"Faculty is not assigned to subject '{subject_code}'"
                )

        if row.grade in {"F", "U", "RA", "AB", "FAIL", "ABSENT"} and row.result_status != "FAIL":
            raise HTTPException(status_code=400, detail=f"Grade '{row.grade}' must have FAIL result status")
        if row.grade in {"SA", "WH", "W", "WITHHELD"} and row.result_status != "WITHHELD":
            raise HTTPException(status_code=400, detail=f"Grade '{row.grade}' must have WITHHELD result status")

        normalized_rows.append((row, student, subject, exam_session))

    imported_count = 0
    for row, student, subject, exam_session in normalized_rows:
        result = db.query(UniversityExamResult).filter(
            UniversityExamResult.student_id == student.id,
            UniversityExamResult.subject_id == subject.id,
            UniversityExamResult.semester == row.semester,
            UniversityExamResult.exam_session == exam_session
        ).first()
        if result is None:
            result = UniversityExamResult(
                student_id=student.id,
                subject_id=subject.id,
                semester=row.semester,
                exam_session=exam_session,
            )
            db.add(result)

        result.grade = row.grade.strip().upper()
        result.result_status = row.result_status
        result.total_marks = row.total_marks
        result.source = "University result PDF"
        imported_count += 1

    try:
        db.commit()
    except Exception:
        db.rollback()
        raise

    record_audit_log(
        db,
        action="IMPORT_UNIVERSITY_RESULTS",
        entity_type="UNIVERSITY_RESULT",
        details=f"Imported {imported_count} university examination result rows for {payload.records[0].exam_session}",
        user_id=current_user.id,
        user_email=current_user.email,
        role=current_user.role
    )

    return {
        "success": True,
        "imported_count": imported_count,
        "message": f"Saved {imported_count} university result records."
    }


@router.get("/university/student/{student_id}")
def get_university_result_history(
    student_id: int,
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user)
):
    target_student = db.query(Student).filter(Student.id == student_id).first()
    if not target_student:
        raise HTTPException(status_code=404, detail="Student not found")

    if current_user.role == "student":
        own_student = db.query(Student).filter(Student.user_id == current_user.id).first()
        if not own_student or own_student.id != student_id:
            raise HTTPException(status_code=403, detail="Unauthorized to view these university results")
    elif current_user.role == "faculty":
        faculty = db.query(Faculty).filter(Faculty.user_id == current_user.id).first()
        if not faculty or target_student.department_id != faculty.department_id:
            raise HTTPException(status_code=403, detail="Faculty can only access result histories in their assigned department")

    results = db.query(UniversityExamResult).filter(
        UniversityExamResult.student_id == student_id
    ).order_by(
        UniversityExamResult.semester.asc(),
        UniversityExamResult.imported_at.asc(),
        UniversityExamResult.id.asc()
    ).all()

    return [
        {
            "id": result.id,
            "student_id": result.student_id,
            "register_number": result.student.register_number,
            "student_name": result.student.full_name,
            "subject_code": result.subject.code,
            "subject_name": result.subject.name,
            "semester": result.semester,
            "exam_session": result.exam_session,
            "grade": result.grade,
            "result_status": result.result_status,
            "total_marks": result.total_marks,
            "imported_at": result.imported_at,
        }
        for result in results
    ]


@router.get("/university/analytics")
def get_university_result_analytics(
    high_risk_arrear_count: int = Query(3, ge=1, le=20),
    db: Session = Depends(get_db),
    current_user: User = Depends(require_role(["admin", "faculty"]))
):
    query = db.query(UniversityExamResult).join(Student).join(Subject)
    if current_user.role == "faculty":
        faculty = db.query(Faculty).filter(Faculty.user_id == current_user.id).first()
        if not faculty:
            raise HTTPException(status_code=403, detail="Faculty profile is required to view university analytics")
        query = query.filter(Student.department_id == faculty.department_id)

    attempts = query.order_by(
        UniversityExamResult.imported_at.asc(),
        UniversityExamResult.id.asc()
    ).all()
    latest_by_course = {}
    session_groups = defaultdict(list)
    failure_attempts = defaultdict(list)

    for attempt in attempts:
        course_key = (attempt.student_id, attempt.subject_id)
        latest_by_course[course_key] = attempt
        session_groups[attempt.exam_session].append(attempt)
        if attempt.result_status == "FAIL":
            failure_attempts[course_key].append(attempt)

    latest_attempts = list(latest_by_course.values())
    current_arrears = [attempt for attempt in latest_attempts if attempt.result_status == "FAIL"]
    student_courses = defaultdict(list)
    subject_courses = defaultdict(list)
    for attempt in latest_attempts:
        student_courses[attempt.student_id].append(attempt)
        subject_courses[attempt.subject_id].append(attempt)

    subject_stats = []
    for subject_id, rows in subject_courses.items():
        subject = rows[0].subject
        failures = sum(row.result_status == "FAIL" for row in rows)
        subject_stats.append({
            "subject_code": subject.code,
            "subject_name": subject.name,
            "semester": subject.semester,
            "total_students": len(rows),
            "pass_percentage": round((len(rows) - failures) / len(rows) * 100, 1),
            "failure_count": failures,
            "current_arrears": failures,
        })
    subject_stats.sort(key=lambda row: (-row["current_arrears"], row["subject_code"]))

    session_trend = []
    for session, rows in session_groups.items():
        passed = sum(row.result_status == "PASS" for row in rows)
        session_trend.append({
            "session": session,
            "pass_percentage": round(passed / len(rows) * 100, 1),
            "pass_count": passed,
            "fail_count": sum(row.result_status == "FAIL" for row in rows),
            "total_attempts": len(rows),
            "imported_at": min(row.imported_at for row in rows),
        })
    session_trend.sort(key=lambda row: (row["imported_at"], row["session"]))

    recurring_arrears = []
    for (student_id, subject_id), failed_rows in failure_attempts.items():
        if len(failed_rows) < 2:
            continue
        latest_failure = failed_rows[-1]
        recurring_arrears.append({
            "register_number": latest_failure.student.register_number,
            "student_name": latest_failure.student.full_name,
            "subject_code": latest_failure.subject.code,
            "subject_name": latest_failure.subject.name,
            "failed_attempts": len(failed_rows),
            "latest_session": latest_failure.exam_session,
        })
    recurring_arrears.sort(key=lambda row: (-row["failed_attempts"], row["register_number"], row["subject_code"]))

    high_risk_students = []
    for student_id, rows in student_courses.items():
        arrears = sum(row.result_status == "FAIL" for row in rows)
        if arrears >= high_risk_arrear_count:
            student = rows[0].student
            high_risk_students.append({
                "register_number": student.register_number,
                "student_name": student.full_name,
                "current_arrears": arrears,
            })
    high_risk_students.sort(key=lambda row: (-row["current_arrears"], row["register_number"]))

    total_students = len(student_courses)
    passed_attempts = sum(row.result_status == "PASS" for row in latest_attempts)
    all_clear_students = sum(
        all(row.result_status != "FAIL" for row in rows)
        for rows in student_courses.values()
    )

    return {
        "summary": {
            "total_students": total_students,
            "all_clear_students": all_clear_students,
            "students_with_arrears": len({row.student_id for row in current_arrears}),
            "total_current_arrears": len(current_arrears),
            "pass_percentage": round(passed_attempts / len(latest_attempts) * 100, 1) if latest_attempts else 0.0,
            "pass_count": passed_attempts,
            "fail_count": len(current_arrears),
            "withheld_count": sum(row.result_status == "WITHHELD" for row in latest_attempts),
            "total_attempts": len(latest_attempts),
        },
        "subject_stats": subject_stats,
        "session_trend": session_trend,
        "recurring_arrears": recurring_arrears,
        "high_risk_students": high_risk_students,
    }
